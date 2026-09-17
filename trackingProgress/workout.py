import json
import os
from copy import copy

import openpyxl
import pandas as pd
import utility


# AA Funzione Per Elaborare I Dati Allenamento E Generare Due DataFrame: workoutData E gymSummary
def processWorkoutData(OPENGYM_DATA, excelData):
    # AA Definizione Dei Tag Colorati Per La Visualizzazione
    tagAllenamento = f"{utility.CLR_ALLENAMENTO}[ALLENAMENTO]{utility.CLR_RESET}"
    
    # AA Caricamento E Parsing Dati Iniziali
    # BB Caricamento Del Catalogo Esercizi Da File Locale (exercisesCatalog.json)
    baseDir = os.path.dirname(os.path.abspath(__file__))
    catalogPath = os.path.join(baseDir, "openGymData", "exercisesCatalog.json")

    exerciseCatalog = {}
    if os.path.exists(catalogPath):
        with open(catalogPath, 'r', encoding='utf-8') as f:
            exerciseCatalog = json.load(f)

    # BB Caricamento File JSON Di openGym
    with open(OPENGYM_DATA, 'r', encoding='utf-8') as f:
        openGymData = json.load(f)

    # CC Integrazione Degli Esercizi Personalizzati Presenti Nel Backup (customEx)
    for cEx in openGymData.get('customEx', []):
        exerciseCatalog[cEx['id']] = cEx.get('n', cEx['id']).title()

    # CC Estrazione Delle Sessioni E Appiattimento Dei Dati Delle Serie
    parsedRows = []
    for w in openGymData.get('workouts', []):
        wDate = w.get('d')
        wName = w.get('name', 'Allenamento')
        wNote = w.get('note', w.get('description', ''))
        startMs = w.get('start')
        endMs = w.get('end')

        for entry in w.get('entries', []):
            exId = entry.get('id')
            exTitle = exerciseCatalog.get(exId, f"Esercizio ({exId})")

            for s in entry.get('sets', []):
                # Escludiamo Eventuali Serie Non Concluse
                if not s.get('done', True):
                    continue

                parsedRows.append({
                    'start_ms': startMs,
                    'end_ms': endMs,
                    'data_str': wDate,
                    'Nome Scheda': wName,
                    'Nome Esercizio': exTitle,
                    'set_type': s.get('phase', 'normal'),
                    'Peso Serie': float(s.get('w', 0)),
                    'Ripetizioni Serie': int(s.get('r', 0)),
                    'Note Post Allenamento': wNote
                })

    workoutData = pd.DataFrame(parsedRows)

    if workoutData.empty:
        return pd.DataFrame(), pd.DataFrame()

    # BB Conversione Date E Calcolo Durata Sessione
    workoutData["Inizio Allenamento"] = pd.to_datetime(workoutData['start_ms'], unit='ms')
    workoutData["Fine Allenamento"] = pd.to_datetime(workoutData["end_ms"], unit='ms')
    workoutData["Data"] = pd.to_datetime(workoutData['data_str']).dt.date
    workoutData["Durata Allenamento"] = (workoutData["Fine Allenamento"] - workoutData["Inizio Allenamento"]).dt.total_seconds() / 60

    # BB Rimozione Colonne Non Necessarie Nel Flusso Iniziale
    columnsToRemove = ['start_ms', 'end_ms', 'data_str', 'Fine Allenamento']
    workoutData = workoutData.drop(columns=columnsToRemove)

    # BB Rimozione Righe Di Riscaldamento E Colonna Tipo Serie
    warmupMask = workoutData['set_type'] == 'warmup'
    rowToRemove = workoutData[warmupMask].index
    workoutData = workoutData.drop(rowToRemove)
    workoutData = workoutData.drop(columns=['set_type'])

    # BB Rimozione Automatica Di Cardio E Attività A Corpo Libero Senza Sovraccarico
    # CC Manteniamo Esclusivamente Le Righe Dove Viene Utilizzato Un Peso Maggiore Di Zero
    # DD Questo Filtro Esclude In Modo Dinamico Corsa, Bici, Plank, Crunch E Corpo Libero Non Zavorrato
    workoutData = workoutData[workoutData['Peso Serie'] > 0]

    # BB Adeguamento Indice Serie In Base 1 Con Ricalcolo Post-Warmup
    # CC Ricalcoliamo La Sequenza Delle Serie Allenanti Usando L'Ora Di Inizio Per Distinguere Le Sessioni Nello Stesso Giorno
    workoutData["Indice Serie"] = workoutData.groupby(['Inizio Allenamento', 'Nome Esercizio']).cumcount() + 1

    # CC Rimozione Della Colonna Di Servizio Temporale Non Più Necessaria Nel DataFrame
    workoutData = workoutData.drop(columns=['Inizio Allenamento'])

    # BB Calcolo Volume E Massimale Per Singola Serie (Espansione Dati Allenamento)
    # CC Calcolo Volume Della Singola Serie (Peso * Reps)
    workoutData["Volume Serie"] = workoutData["Peso Serie"] * workoutData["Ripetizioni Serie"]

    # CC Calcolo Massimale Stimato Di Riga Tramite Formula Epley
    # DD Formula Di Epley: Peso * (1 + Reps / 30)
    workoutData["Massimale Stimato Serie"] = workoutData["Peso Serie"] * (1 + workoutData["Ripetizioni Serie"] / 30)

    # CC Arrotondamento E Conversione In Intero Delle Nuove Metriche Di Riga
    # EE La Conversione In Int Non Genera Errori Poiché Non Vi Sono Più Valori NaN Nel Dataset
    workoutData["Volume Serie"] = workoutData["Volume Serie"].round().astype(int)
    workoutData["Massimale Stimato Serie"] = workoutData["Massimale Stimato Serie"].round().astype(int)


    # AA Setting Colonne DataFrame Per "Dati Palestra" (Bacheca Record Assoluti)
    # BB Copia Del DataFrame Master Espanso
    gymSummary = workoutData.copy()

    # BB Raggruppamento Temporaneo Per Data E Nome Esercizio (85 Sessioni)
    groupedSummary = gymSummary.groupby(['Data', 'Nome Esercizio'])

    gymSummary = groupedSummary.agg({
        'Indice Serie': 'count',
        'Ripetizioni Serie': 'mean',
        'Peso Serie': 'mean',
        'Volume Serie': 'sum',
        'Massimale Stimato Serie': 'max',
        'Note Post Allenamento': 'first'
    }).reset_index()

    # BB Rinomina Delle Metriche Aggregate Di Sessione
    renameMapGym = {
        'Indice Serie': 'Numero Serie',
        'Ripetizioni Serie': 'Ripetizioni Medie',
        'Peso Serie': 'Peso Medio',
        'Volume Serie': 'Volume Totale (Kg)',
        'Massimale Stimato Serie': 'Massimale Stimato (1RM)'
    }
    gymSummary = gymSummary.rename(columns=renameMapGym)

    # BB Approssimazione Medie E Arrotondamenti A Numeri Interi
    gymSummary['Ripetizioni Medie'] = gymSummary['Ripetizioni Medie'].round().astype(int)
    gymSummary['Peso Medio'] = gymSummary['Peso Medio'].round().astype(int)

    # BB Creazione Colonna Schema Riassuntivo Di Sessione
    # CC Utilizzo Di Una Funzione Lambda Per Formattare I Tre Parametri Chiave
    gymSummary['Schema'] = gymSummary.apply(lambda row: f"{row['Numero Serie']}x{row['Ripetizioni Medie']}x{row['Peso Medio']}kg", axis=1)

    # BB Riduzione A Un Record Assoluto Storico Per Ciascun Esercizio 
    # CC Identificazione Della Riga Con Il Massimo 1RM Storico Per Ogni Singolo Gruppo Esercizio
    gymSummary = gymSummary.loc[gymSummary.groupby('Nome Esercizio')['Massimale Stimato (1RM)'].idxmax()].reset_index(drop=True)

    # CC Ordinamento Alfabetico Degli Esercizi Per Una Migliore Resa Estetica Su Excel
    gymSummary = gymSummary.sort_values(by='Nome Esercizio').reset_index(drop=True)
    # DD -----------------------------------------------

    # EE -----------------------------------------------
    # AA Confronto Con Excel E Filtraggio Dati Allenamento Già Presenti
    existingWorkout = excelData['Dati Allenamento']

    # CC Definizione Delle Colonne Per Entrambi I DataFrame Per Garantire Coerenza Di Scrittura
    columnsWorkout = ['Data', 'Durata Allenamento', 'Nome Scheda', 'Nome Esercizio', 
                    'Indice Serie', 'Ripetizioni Serie', 'Peso Serie', 'Volume Serie', 
                    'Massimale Stimato Serie', 'Note Post Allenamento']

    columnsGym = ['Nome Esercizio', 'Schema', 'Volume Totale (Kg)', 'Data', 'Numero Serie', 
                'Ripetizioni Medie', 'Peso Medio', 'Massimale Stimato (1RM)', 'Note Post Allenamento']

    if existingWorkout.empty:
        print(f"\n{tagAllenamento} Nessun Dato Allenamento Precedente Trovato In Excel. Tutti I Dati Saranno Importati.\n")
    else:
        # CC Filtraggio Dati Allenamento Già Presenti In Excel Tramite Calcolo Della Data Massima
        existingWorkoutDate = pd.to_datetime(existingWorkout['Data']).dt.date
        lastExistingDate = existingWorkoutDate.max()
        print(f"\n{tagAllenamento} Ultima Data Allenamento Presente In Excel: {utility.CLR_BOLD}{lastExistingDate}{utility.CLR_RESET}\n")
        lastDateMask = workoutData['Data'] > lastExistingDate
        workoutData = workoutData[lastDateMask]

    # CC Riordino E Selezione Delle Colonne Eseguito In Ogni Caso (Sia Con Excel Vuoto Che Popolato)
    workoutData = workoutData[columnsWorkout]
    gymSummary = gymSummary[columnsGym]
    # EE -----------------------------------------------

    return workoutData, gymSummary

# AA Funzione Per Aggiornare Scheda Dati Allenamento In Excel Con I Nuovi Allenamenti
def updateTrainData(workoutData, excelData, EXCEL_PATH):
    # BB Definizione Dei Tag Colorati Per La Visualizzazione
    tagAllenamento = f"{utility.CLR_ALLENAMENTO}[ALLENAMENTO]{utility.CLR_RESET}"

    # BB Controllo Preliminare Se Ci Sono Nuovi Record Di Allenamento Rilevati
    if workoutData is None or len(workoutData) == 0:
        print(f"{tagAllenamento} Nessun Nuovo Allenamento Rilevato Da Inserire.\n")
        return
    else:
        print(f"{tagAllenamento} Rilevati {utility.CLR_BOLD}{len(workoutData)}{utility.CLR_RESET} Nuovi Record Da Salvare.\n")

    # CC Preparazione Dei Dati E Allineamento Colonne Per La Scheda Dati Allenamento
    dfToAppend = workoutData.copy()
    if 'Data Allenamento' in dfToAppend.columns and 'Data' not in dfToAppend.columns:
        dfToAppend = dfToAppend.rename(columns={'Data Allenamento': 'Data'})

    excelColumns = [
        'Data', 'Durata Allenamento', 'Nome Scheda', 'Nome Esercizio', 
        'Indice Serie', 'Ripetizioni Serie', 'Peso Serie', 'Volume Serie', 
        'Massimale Stimato Serie', 'Note Post Allenamento'
    ]
    dfToAppend = dfToAppend[excelColumns]

    # DD Caricamento Del File Excel Con Openpyxl Per Preservare Stili, Colori E Larghezze Delle Celle E Aggiunta Nuove Righe
    excelFile = openpyxl.load_workbook(EXCEL_PATH)
    workoutSheet = excelFile['Dati Allenamento']

    # DD Scrittura Delle Nuove Righe Nel Foglio Di Lavoro
    for idx, row in dfToAppend.iterrows():
        # EE Creazione Della Lista Ordinata Dei Nuovi Valori Per Ciascuna Riga
        newValue = []
        for colName in excelColumns:
            val = row[colName]
            
            # Gestione Dei Valori Nulli O Incompatibili Con Excel
            if pd.isna(val):
                val = ""
                
            # Conversione Del Formato Data Di Pandas Per Evitare Errori Di Scrittura
            if colName == 'Data' and isinstance(val, pd.Timestamp):
                val = val.date()
                
            newValue.append(val)
            
        workoutSheet.append(newValue)
        lastRowIdx = workoutSheet.max_row
        
        # EE Copia Della Formattazione Dalla Riga Precedente (Incluso Formato Data, Carattere E Allineamenti)
        if lastRowIdx > 2: # Riga 1 = Intestazioni, Riga 2 = Primo Dato (Se Esiste, Copiamo Da Lì)
            for col_idx in range(1, len(newValue) + 1):
                sourceCell = workoutSheet.cell(row=lastRowIdx - 1, column=col_idx)
                targetCell = workoutSheet.cell(row=lastRowIdx, column=col_idx)
                
                # Copiamo Gli Stili E Il Formato Numerico Se Presenti (Silenziando I Falsi Positivi Di Pylance)
                if sourceCell.has_style:
                    targetCell.font = copy(sourceCell.font)  # type: ignore
                    targetCell.border = copy(sourceCell.border)  # type: ignore
                    targetCell.fill = copy(sourceCell.fill)  # type: ignore
                    targetCell.alignment = copy(sourceCell.alignment)  # type: ignore
                    targetCell.number_format = sourceCell.number_format

        else:
            # Formattazione Manuale Solo Se È La Prima Riga Di Dati In Assoluto
            # Formato Data Per La Colonna 1 (Data Allenamento)
            workoutSheet.cell(row=lastRowIdx, column=1).number_format = 'DD-MM-YYYY'
            
    # DD Salvataggio Del File Excel Per Applicare Le Modifiche Senza Ricreare I Fogli
    excelFile.save(EXCEL_PATH)
    excelFile.close()

    # DD Sincronizzazione Del DataFrame In Memoria Per Coerenza Della Sessione
    excelData['Dati Allenamento'] = pd.concat([excelData['Dati Allenamento'], dfToAppend], ignore_index=True)

    print(f"{tagAllenamento} {utility.CLR_BOLD}Dati Di Allenamento E Calcoli Salvati Con Successo Nel File Excel!{utility.CLR_RESET}\n")


# AA Funzione Per Ricreare La Scheda Dati Palestra In Excel Con I Nuovi Record
def updateGymData(gymSummary, excelData, EXCEL_PATH):
    # BB Definizione Dei Tag Colorati Per La Visualizzazione
    tagPalestra = f"{utility.CLR_PALESTRA}[PALESTRA]{utility.CLR_RESET}"

    # BB Controllo Preliminare Se Ci Sono Record Di Palestra Da Aggiornare
    if gymSummary is None or len(gymSummary) == 0:
        print(f"{tagPalestra} Nessun Record Rilevato Da Aggiornare.\n")
        return
    else:
        print(f"{tagPalestra} Aggiornamento Di {utility.CLR_BOLD}{len(gymSummary)}{utility.CLR_RESET} Record Storici In Corso...\n")

    # BB Preparazione Dei Dati E Allineamento Colonne Per La Scheda Dati Palestra
    dfToOverwrite = gymSummary.copy()
    
    excelColumns = [
        'Nome Esercizio', 'Schema', 'Volume Totale (Kg)', 'Data', 
        'Numero Serie', 'Ripetizioni Medie', 'Peso Medio', 
        'Massimale Stimato (1RM)', 'Note Post Allenamento'
    ]
    dfToOverwrite = dfToOverwrite[excelColumns]

    # CC Caricamento Del File Excel Con Openpyxl Per Preservare Stili, Colori E Larghezze Delle Celle E Sovrascrittura
    excelFile = openpyxl.load_workbook(EXCEL_PATH)
    gymSheet = excelFile['Dati Palestra']

    # CC Salvataggio Degli Stili Esistenti Dalla Riga 2 Come Modello Per Il Popolamento
    styleTemplates = {}
    if gymSheet.max_row >= 2:
        for col_idx in range(1, len(excelColumns) + 1):
            cell = gymSheet.cell(row=2, column=col_idx)
            styleTemplates[col_idx] = {
                'font': copy(cell.font),
                'border': copy(cell.border),
                'fill': copy(cell.fill),
                'alignment': copy(cell.alignment),
                'number_format': cell.number_format
            }

    # CC Rimozione Di Tutti I Record Esistenti Dalla Scheda Per Consentire La Sovrascrittura Completa
    if gymSheet.max_row > 1:
        gymSheet.delete_rows(2, amount=gymSheet.max_row)

    # CC Scrittura Dei Nuovi Record Nel Foglio Di Lavoro
    for idx, row in dfToOverwrite.iterrows():
        # CC Creazione Della Lista Ordinata Dei Nuovi Valori Per Ciascuna Riga
        newValue = []
        for colName in excelColumns:
            val = row[colName]
            
            # EE Gestione Dei Valori Nulli O Incompatibili Con Excel
            if pd.isna(val):
                val = ""
                
            # EE Conversione Del Formato Data Di Pandas Per Evitare Errori Di Scrittura
            if colName == 'Data' and isinstance(val, pd.Timestamp):
                val = val.date()
                
            newValue.append(val)
            
        gymSheet.append(newValue)
        lastRowIdx = gymSheet.max_row
        
        # DD Copia Della Formattazione Tramite I Modelli Salvati In Precedenza
        if len(styleTemplates) > 0:
            for col_idx in range(1, len(newValue) + 1):
                targetCell = gymSheet.cell(row=lastRowIdx, column=col_idx)
                tpl = styleTemplates[col_idx]
                
                # Copiamo Gli Stili E Il Formato Numerico (Silenziando I Falsi Positivi Di Pylance)
                targetCell.font = copy(tpl['font'])  # type: ignore
                targetCell.border = copy(tpl['border'])  # type: ignore
                targetCell.fill = copy(tpl['fill'])  # type: ignore
                targetCell.alignment = copy(tpl['alignment'])  # type: ignore
                targetCell.number_format = tpl['number_format']
        else:
            # EE Formattazione Manuale Solo Se La Scheda Era Completamente Vuota
            # Formato Data Per La Colonna 4 (Data)
            gymSheet.cell(row=lastRowIdx, column=4).number_format = 'DD-MM-YYYY'

    # CC Salvataggio Del File Excel Per Applicare Le Modifiche Senza Ricreare I Fogli
    excelFile.save(EXCEL_PATH)
    excelFile.close()

    # CC Sincronizzazione Del DataFrame In Memoria Per Coerenza Della Sessione
    excelData['Dati Palestra'] = dfToOverwrite

    print(f"{tagPalestra} {utility.CLR_BOLD}Record Storici Personali Aggiornati E Salvati Con Successo!{utility.CLR_RESET}\n")